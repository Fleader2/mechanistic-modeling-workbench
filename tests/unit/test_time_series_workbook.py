"""Tests for ``app.time_series.template``/``app.time_series.parser`` -- pure logic, no
database, no model, no Antimony species name involved anywhere."""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models.enums import MeasurementValueStatus
from app.time_series.contract import DATA_SHEET_NAME, METADATA_SHEET_NAME, TEMPLATE_VERSION
from app.time_series.parser import WorkbookMalformedError, parse_workbook
from tests.time_series_helpers import (
    FIXTURE_HEADERS,
    FIXTURE_ROWS,
    blank_template_bytes,
    declare_independent_variables,
    declare_measurement_uncertainty,
    fixture_workbook_bytes,
    populate_data_tab,
)

pytestmark = pytest.mark.unit


class TestTemplate:
    def test_generates_exactly_two_named_sheets_in_order(self):
        workbook = load_workbook(BytesIO(blank_template_bytes()))
        assert workbook.sheetnames == [METADATA_SHEET_NAME, DATA_SHEET_NAME]

    def test_metadata_sheet_prefills_the_template_version_and_nothing_else(self):
        workbook = load_workbook(BytesIO(blank_template_bytes()))
        sheet = workbook[METADATA_SHEET_NAME]
        values = {
            sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
            for r in range(2, sheet.max_row + 1)
            if isinstance(sheet.cell(row=r, column=1).value, str)
        }
        assert values.get("Template version") == TEMPLATE_VERSION
        assert values.get("Dataset name") is None
        assert values.get("Organism") is None

    def test_data_sheet_has_the_five_structural_headers_and_nothing_else(self):
        workbook = load_workbook(BytesIO(blank_template_bytes()))
        sheet = workbook[DATA_SHEET_NAME]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        assert header_row == (
            "sample_id",
            "condition_id",
            "biological_replicate",
            "technical_replicate",
            "time",
        )


class TestMalformedWorkbooks:
    def test_unreadable_bytes_are_rejected(self):
        with pytest.raises(WorkbookMalformedError) as exc_info:
            parse_workbook(b"this is not a zip file at all")
        assert [f.code for f in exc_info.value.findings] == ["WORKBOOK_UNREADABLE"]

    def test_missing_data_sheet_is_rejected(self):
        workbook = load_workbook(BytesIO(blank_template_bytes()))
        del workbook[DATA_SHEET_NAME]
        buffer = BytesIO()
        workbook.save(buffer)
        with pytest.raises(WorkbookMalformedError) as exc_info:
            parse_workbook(buffer.getvalue())
        assert [f.code for f in exc_info.value.findings] == ["MISSING_REQUIRED_SHEET"]

    def test_missing_metadata_sheet_is_rejected(self):
        workbook = load_workbook(BytesIO(blank_template_bytes()))
        del workbook[METADATA_SHEET_NAME]
        buffer = BytesIO()
        workbook.save(buffer)
        with pytest.raises(WorkbookMalformedError) as exc_info:
            parse_workbook(buffer.getvalue())
        assert [f.code for f in exc_info.value.findings] == ["MISSING_REQUIRED_SHEET"]

    def test_blank_template_with_no_measurement_columns_is_rejected(self):
        with pytest.raises(WorkbookMalformedError) as exc_info:
            parse_workbook(blank_template_bytes())
        assert [f.code for f in exc_info.value.findings] == ["NO_MEASUREMENT_COLUMNS"]

    def test_missing_time_column_is_rejected(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("sample_id", "glucose_mM"),
            rows=((None, 1.0),),
        )
        with pytest.raises(WorkbookMalformedError) as exc_info:
            parse_workbook(data)
        assert [f.code for f in exc_info.value.findings] == ["MISSING_TIME_COLUMN"]

    def test_duplicate_measurement_column_header_is_rejected(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "glucose_mM"),
            rows=((0, 1.0, 2.0),),
        )
        with pytest.raises(WorkbookMalformedError) as exc_info:
            parse_workbook(data)
        assert [f.code for f in exc_info.value.findings] == ["DUPLICATE_DATA_COLUMN_HEADER"]

    def test_duplicate_structural_column_header_is_rejected(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "Time", "glucose_mM"),
            rows=((0, 0, 1.0),),
        )
        with pytest.raises(WorkbookMalformedError) as exc_info:
            parse_workbook(data)
        assert [f.code for f in exc_info.value.findings] == ["DUPLICATE_DATA_COLUMN_HEADER"]

    def test_time_column_match_is_case_insensitive(self):
        # Confirms the fixture's own "Time" header (capitalized) is recognized -- never
        # rejected as MISSING_TIME_COLUMN.
        data = populate_data_tab(
            blank_template_bytes(), headers=("TIME", "glucose_mM"), rows=((0, 1.0),)
        )
        parsed = parse_workbook(data)
        assert parsed.rows[0].time_value == 0.0


class TestValueClassification:
    def test_missing_unknown_not_applicable_and_a_real_zero_are_all_distinguished(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM"),
            rows=(
                (0, None),
                (10, "NA"),
                (20, "unknown"),
                (30, "not a number"),
                (40, 0),
            ),
        )
        parsed = parse_workbook(data)
        statuses = {row.row_index: row.measurements["glucose_mM"] for row in parsed.rows}

        assert statuses[0].status == MeasurementValueStatus.MISSING
        assert statuses[0].value is None

        assert statuses[1].status == MeasurementValueStatus.NOT_APPLICABLE
        assert statuses[1].value is None

        assert statuses[2].status == MeasurementValueStatus.UNKNOWN
        assert statuses[2].value is None

        assert statuses[3].status == MeasurementValueStatus.UNRECOGNIZED
        assert statuses[3].raw == "not a number"

        # The regression case: a real, reported zero is MEASURED, never collapsed to MISSING.
        assert statuses[4].status == MeasurementValueStatus.MEASURED
        assert statuses[4].value == 0.0

        unrecognized_findings = [
            f for f in parsed.findings if f.code == "DATA_CELL_UNRECOGNIZED_VALUE"
        ]
        assert len(unrecognized_findings) == 1
        assert unrecognized_findings[0].row_index == 3

    def test_na_and_unknown_tokens_are_case_insensitive(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM"),
            rows=((0, "n/a"), (10, "UNK")),
        )
        parsed = parse_workbook(data)
        row0 = parsed.rows[0].measurements["glucose_mM"]
        row1 = parsed.rows[1].measurements["glucose_mM"]
        assert row0.status == MeasurementValueStatus.NOT_APPLICABLE
        assert row1.status == MeasurementValueStatus.UNKNOWN

    def test_a_boolean_cell_is_never_silently_coerced_to_zero_or_one(self):
        data = populate_data_tab(
            blank_template_bytes(), headers=("time", "glucose_mM"), rows=((0, True),)
        )
        parsed = parse_workbook(data)
        value = parsed.rows[0].measurements["glucose_mM"]
        assert value.status == MeasurementValueStatus.UNRECOGNIZED


class TestMeasurementColumnDeclarations:
    def test_a_declared_column_picks_up_its_unit_and_description(self):
        data = populate_data_tab(
            blank_template_bytes(), headers=("time", "glucose_mM"), rows=((0, 1.0),)
        )
        workbook = load_workbook(BytesIO(data))
        metadata_sheet = workbook["Metadata"]
        section_row = next(
            r
            for r in range(1, metadata_sheet.max_row + 1)
            if isinstance(metadata_sheet.cell(row=r, column=1).value, str)
            and metadata_sheet.cell(row=r, column=1).value.startswith(
                "## Measurement column definitions"
            )
        )
        header_row = section_row + 1
        metadata_sheet.cell(row=header_row + 1, column=1, value="glucose_mM")
        metadata_sheet.cell(row=header_row + 1, column=2, value="mM")
        metadata_sheet.cell(row=header_row + 1, column=3, value="Extracellular glucose")
        buffer = BytesIO()
        workbook.save(buffer)

        parsed = parse_workbook(buffer.getvalue())
        column = next(c for c in parsed.measurement_columns if c.column_name == "glucose_mM")
        assert column.unit == "mM"
        assert column.description == "Extracellular glucose"
        assert column.declared_in_metadata is True
        assert not any(f.code == "MEASUREMENT_COLUMN_UNIT_UNDECLARED" for f in parsed.findings)

    def test_an_undeclared_column_is_flagged_but_still_imported(self):
        data = populate_data_tab(
            blank_template_bytes(), headers=("time", "glucose_mM"), rows=((0, 1.0),)
        )
        parsed = parse_workbook(data)
        column = parsed.measurement_columns[0]
        assert column.declared_in_metadata is False
        assert column.unit is None
        assert any(
            f.code == "MEASUREMENT_COLUMN_UNIT_UNDECLARED" and f.column_name == "glucose_mM"
            for f in parsed.findings
        )


class TestIndependentVariableColumns:
    """Joint Calibration Workflow prerequisite increment (docs/29 §4, docs/31)."""

    def test_a_workbook_with_no_independent_variable_declarations_is_unchanged(self):
        """Preserves every existing import byte-for-byte: with no "Independent variable
        definitions" sub-table at all, every non-structural column is still a measurement
        column, exactly as before this increment existed."""
        data = populate_data_tab(
            blank_template_bytes(), headers=("time", "glucose_mM"), rows=((0, 1.0),)
        )
        parsed = parse_workbook(data)
        assert parsed.independent_variable_columns == []
        assert [c.column_name for c in parsed.measurement_columns] == ["glucose_mM"]
        assert parsed.rows[0].independent_variables == {}

    def test_a_declared_independent_variable_column_is_excluded_from_measurement_columns(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "dose_uM", "glucose_mM"),
            rows=((0, 10.0, 1.0),),
        )
        data = declare_independent_variables(
            data, {"dose_uM": ("uM", "Drug concentration applied at t=0")}
        )
        parsed = parse_workbook(data)

        assert [c.column_name for c in parsed.independent_variable_columns] == ["dose_uM"]
        iv_column = parsed.independent_variable_columns[0]
        assert iv_column.unit == "uM"
        assert iv_column.description == "Drug concentration applied at t=0"

        # Never also a measurement column, and never flagged as an undeclared measurement.
        assert [c.column_name for c in parsed.measurement_columns] == ["glucose_mM"]
        assert not any(f.column_name == "dose_uM" for f in parsed.findings)

        assert parsed.rows[0].independent_variables["dose_uM"].value == 10.0
        assert "dose_uM" not in parsed.rows[0].measurements

    def test_an_independent_variable_cell_preserves_real_zero_and_discloses_unrecognized(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "dose_uM", "glucose_mM"),
            rows=((0, 0.0, 1.0), (60, "garbled", 2.0)),
        )
        data = declare_independent_variables(data, {"dose_uM": ("uM", None)})
        parsed = parse_workbook(data)

        assert parsed.rows[0].independent_variables["dose_uM"].value == 0.0
        assert parsed.rows[0].independent_variables["dose_uM"].status == "MEASURED"
        assert parsed.rows[1].independent_variables["dose_uM"].status == "UNRECOGNIZED"
        assert any(
            f.code == "DATA_CELL_UNRECOGNIZED_VALUE" and f.column_name == "dose_uM"
            for f in parsed.findings
        )

    def test_an_orphan_independent_variable_declaration_is_disclosed(self):
        data = populate_data_tab(
            blank_template_bytes(), headers=("time", "glucose_mM"), rows=((0, 1.0),)
        )
        data = declare_independent_variables(data, {"dose_uM": ("uM", None)})
        parsed = parse_workbook(data)
        assert parsed.independent_variable_columns == []
        assert any(
            f.code == "ORPHAN_INDEPENDENT_VARIABLE_COLUMN_DEFINITION"
            and f.column_name == "dose_uM"
            for f in parsed.findings
        )


class TestMeasurementUncertaintyColumns:
    """Calibration configuration close-out increment (docs/29 §5, docs/32)."""

    def test_a_workbook_with_no_uncertainty_declarations_is_unchanged(self):
        data = populate_data_tab(
            blank_template_bytes(), headers=("time", "glucose_mM"), rows=((0, 1.0),)
        )
        parsed = parse_workbook(data)
        assert parsed.measurement_uncertainty_columns == []
        assert [c.column_name for c in parsed.measurement_columns] == ["glucose_mM"]
        assert parsed.rows[0].measurement_uncertainty == {}

    def test_a_declared_uncertainty_column_is_excluded_from_measurement_columns(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "glucose_mM__sigma"),
            rows=((0, 1.0, 0.1),),
        )
        data = declare_measurement_uncertainty(
            data, {"glucose_mM__sigma": ("glucose_mM", "SD", "mM", "Standard deviation")}
        )
        parsed = parse_workbook(data)

        assert [c.column_name for c in parsed.measurement_uncertainty_columns] == [
            "glucose_mM__sigma"
        ]
        uncertainty_column = parsed.measurement_uncertainty_columns[0]
        assert uncertainty_column.measurement_column_name == "glucose_mM"
        assert uncertainty_column.statistic_kind == "STANDARD_DEVIATION"
        assert uncertainty_column.unit == "mM"
        assert uncertainty_column.description == "Standard deviation"

        assert [c.column_name for c in parsed.measurement_columns] == ["glucose_mM"]
        assert not any(f.column_name == "glucose_mM__sigma" for f in parsed.findings)

        # Keyed by the *measurement* column name, never the uncertainty column's own name.
        assert parsed.rows[0].measurement_uncertainty["glucose_mM"].value == 0.1
        assert "glucose_mM__sigma" not in parsed.rows[0].measurements
        assert "glucose_mM__sigma" not in parsed.rows[0].measurement_uncertainty

    def test_an_uncertainty_cell_preserves_real_zero_and_discloses_unrecognized(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "glucose_mM__sigma"),
            rows=((0, 1.0, 0.0), (60, 2.0, "garbled")),
        )
        data = declare_measurement_uncertainty(
            data, {"glucose_mM__sigma": ("glucose_mM", "SD", None, None)}
        )
        parsed = parse_workbook(data)

        assert parsed.rows[0].measurement_uncertainty["glucose_mM"].value == 0.0
        assert parsed.rows[0].measurement_uncertainty["glucose_mM"].status == "MEASURED"
        assert parsed.rows[1].measurement_uncertainty["glucose_mM"].status == "UNRECOGNIZED"
        assert any(
            f.code == "DATA_CELL_UNRECOGNIZED_VALUE" and f.column_name == "glucose_mM"
            for f in parsed.findings
        )

    def test_a_negative_uncertainty_value_is_preserved_and_disclosed(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "glucose_mM__sigma"),
            rows=((0, 1.0, -0.5),),
        )
        data = declare_measurement_uncertainty(
            data, {"glucose_mM__sigma": ("glucose_mM", "SD", None, None)}
        )
        parsed = parse_workbook(data)

        assert parsed.rows[0].measurement_uncertainty["glucose_mM"].value == -0.5
        assert any(
            f.code == "NEGATIVE_UNCERTAINTY_VALUE" and f.column_name == "glucose_mM"
            for f in parsed.findings
        )

    def test_an_uncertainty_column_with_no_declared_target_is_disclosed_and_unused(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "mystery_sigma"),
            rows=((0, 1.0, 0.1),),
        )
        data = declare_measurement_uncertainty(data, {"mystery_sigma": ("", "SD", None, None)})
        parsed = parse_workbook(data)

        assert parsed.measurement_uncertainty_columns == []
        assert parsed.rows[0].measurement_uncertainty == {}
        assert any(
            f.code == "MEASUREMENT_UNCERTAINTY_COLUMN_MISSING_TARGET"
            and f.column_name == "mystery_sigma"
            for f in parsed.findings
        )

    def test_an_uncertainty_column_naming_a_nonexistent_measurement_column_is_disclosed(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "lactate_mM__sigma"),
            rows=((0, 1.0, 0.1),),
        )
        data = declare_measurement_uncertainty(
            data, {"lactate_mM__sigma": ("lactate_mM", "SD", None, None)}
        )
        parsed = parse_workbook(data)

        assert parsed.measurement_uncertainty_columns == []
        assert any(
            f.code == "MEASUREMENT_UNCERTAINTY_COLUMN_TARGET_NOT_FOUND"
            and f.column_name == "lactate_mM__sigma"
            for f in parsed.findings
        )

    def test_two_uncertainty_columns_targeting_the_same_measurement_column_the_first_wins(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "sigma_a", "sigma_b"),
            rows=((0, 1.0, 0.1, 0.2),),
        )
        data = declare_measurement_uncertainty(
            data,
            {
                "sigma_a": ("glucose_mM", "SD", None, None),
                "sigma_b": ("glucose_mM", "SD", None, None),
            },
        )
        parsed = parse_workbook(data)

        assert [c.column_name for c in parsed.measurement_uncertainty_columns] == ["sigma_a"]
        assert parsed.rows[0].measurement_uncertainty["glucose_mM"].value == 0.1
        assert any(
            f.code == "DUPLICATE_MEASUREMENT_UNCERTAINTY_TARGET" and f.column_name == "sigma_b"
            for f in parsed.findings
        )

    def test_an_orphan_uncertainty_declaration_is_disclosed(self):
        data = populate_data_tab(
            blank_template_bytes(), headers=("time", "glucose_mM"), rows=((0, 1.0),)
        )
        data = declare_measurement_uncertainty(
            data, {"glucose_mM__sigma": ("glucose_mM", "SD", None, None)}
        )
        parsed = parse_workbook(data)
        assert parsed.measurement_uncertainty_columns == []
        assert any(
            f.code == "ORPHAN_MEASUREMENT_UNCERTAINTY_COLUMN_DEFINITION"
            and f.column_name == "glucose_mM__sigma"
            for f in parsed.findings
        )

    def test_a_column_declared_both_independent_variable_and_uncertainty_prefers_independent(
        self,
    ):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "ambiguous_col"),
            rows=((0, 1.0, 5.0),),
        )
        data = declare_independent_variables(data, {"ambiguous_col": (None, None)})
        data = declare_measurement_uncertainty(
            data, {"ambiguous_col": ("glucose_mM", "SD", None, None)}
        )
        parsed = parse_workbook(data)

        assert [c.column_name for c in parsed.independent_variable_columns] == ["ambiguous_col"]
        assert parsed.measurement_uncertainty_columns == []
        assert parsed.rows[0].independent_variables["ambiguous_col"].value == 5.0
        assert any(
            f.code == "COLUMN_DECLARED_BOTH_INDEPENDENT_VARIABLE_AND_UNCERTAINTY"
            and f.column_name == "ambiguous_col"
            for f in parsed.findings
        )

    def test_an_uncertainty_column_with_a_blank_statistic_is_disclosed_and_unused(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "glucose_mM__sigma"),
            rows=((0, 1.0, 0.1),),
        )
        data = declare_measurement_uncertainty(
            data, {"glucose_mM__sigma": ("glucose_mM", None, None, None)}
        )
        parsed = parse_workbook(data)

        assert parsed.measurement_uncertainty_columns == []
        assert parsed.rows[0].measurement_uncertainty == {}
        assert any(
            f.code == "MEASUREMENT_UNCERTAINTY_COLUMN_STATISTIC_KIND_UNRECOGNIZED"
            and f.column_name == "glucose_mM__sigma"
            for f in parsed.findings
        )

    def test_an_uncertainty_column_with_an_unrecognized_statistic_is_disclosed_and_unused(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "glucose_mM__sigma"),
            rows=((0, 1.0, 0.1),),
        )
        data = declare_measurement_uncertainty(
            data, {"glucose_mM__sigma": ("glucose_mM", "a spread of some kind", None, None)}
        )
        parsed = parse_workbook(data)

        assert parsed.measurement_uncertainty_columns == []
        assert any(
            f.code == "MEASUREMENT_UNCERTAINTY_COLUMN_STATISTIC_KIND_UNRECOGNIZED"
            and f.column_name == "glucose_mM__sigma"
            for f in parsed.findings
        )

    def test_variance_and_standard_error_statistic_kinds_are_each_recognized(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "var_col"),
            rows=((0, 1.0, 0.1),),
        )
        data = declare_measurement_uncertainty(
            data, {"var_col": ("glucose_mM", "variance", None, None)}
        )
        parsed = parse_workbook(data)
        assert [c.statistic_kind for c in parsed.measurement_uncertainty_columns] == ["VARIANCE"]

        data2 = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "se_col"),
            rows=((0, 1.0, 0.2),),
        )
        data2 = declare_measurement_uncertainty(
            data2, {"se_col": ("glucose_mM", "Standard Error", None, None)}
        )
        parsed2 = parse_workbook(data2)
        assert [c.statistic_kind for c in parsed2.measurement_uncertainty_columns] == [
            "STANDARD_ERROR"
        ]

    def test_a_non_finite_measurement_cell_is_unrecognized_never_measured(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM"),
            rows=((0, "inf"), (60, "-inf"), (120, "nan")),
        )
        parsed = parse_workbook(data)
        for row in parsed.rows:
            assert row.measurements["glucose_mM"].status == "UNRECOGNIZED"
            assert row.measurements["glucose_mM"].value is None

    def test_a_non_finite_uncertainty_cell_is_unrecognized_never_measured(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "glucose_mM", "glucose_mM__sigma"),
            rows=((0, 1.0, "inf"),),
        )
        data = declare_measurement_uncertainty(
            data, {"glucose_mM__sigma": ("glucose_mM", "SD", None, None)}
        )
        parsed = parse_workbook(data)
        assert parsed.rows[0].measurement_uncertainty["glucose_mM"].status == "UNRECOGNIZED"
        assert parsed.rows[0].measurement_uncertainty["glucose_mM"].value is None


class TestConditionLegend:
    def test_a_filled_condition_legend_row_is_parsed(self):
        data = populate_data_tab(
            blank_template_bytes(),
            headers=("time", "condition_id", "glucose_mM"),
            rows=((0, "ctrl", 1.0),),
        )
        workbook = load_workbook(BytesIO(data))
        metadata_sheet = workbook["Metadata"]
        section_row = next(
            r
            for r in range(1, metadata_sheet.max_row + 1)
            if isinstance(metadata_sheet.cell(row=r, column=1).value, str)
            and metadata_sheet.cell(row=r, column=1).value.startswith(
                "## Condition / treatment legend"
            )
        )
        header_row = section_row + 1
        metadata_sheet.cell(row=header_row + 1, column=1, value="ctrl")
        metadata_sheet.cell(row=header_row + 1, column=2, value="Vehicle control")
        buffer = BytesIO()
        workbook.save(buffer)

        parsed = parse_workbook(buffer.getvalue())
        assert parsed.condition_legend == [("ctrl", "Vehicle control")]


class TestSuppliedRegressionFixture:
    """The exact literal fixture supplied with this increment's own task instructions."""

    def test_round_trips_without_loss(self):
        parsed = parse_workbook(fixture_workbook_bytes())

        assert [c.column_name for c in parsed.measurement_columns] == list(FIXTURE_HEADERS[1:])
        assert len(parsed.rows) == len(FIXTURE_ROWS)

        for row, fixture_row in zip(parsed.rows, FIXTURE_ROWS, strict=True):
            assert row.time_value == float(fixture_row[0])
            for column_name, expected in zip(FIXTURE_HEADERS[1:], fixture_row[1:], strict=True):
                measurement = row.measurements[column_name]
                assert measurement.status == MeasurementValueStatus.MEASURED
                assert measurement.value == pytest.approx(float(expected))

    def test_the_reported_zero_at_t_120_for_pmektot_is_preserved_as_measured(self):
        parsed = parse_workbook(fixture_workbook_bytes())
        row_t120 = next(r for r in parsed.rows if r.time_value == 120.0)
        value = row_t120.measurements["[pMEKtot]"]
        assert value.status == MeasurementValueStatus.MEASURED
        assert value.value == 0.0

    def test_every_blank_required_metadata_field_is_disclosed_as_a_finding_never_fabricated(self):
        parsed = parse_workbook(fixture_workbook_bytes())
        assert parsed.metadata["organism"] is None
        assert parsed.metadata["experimental_date"] is None
        missing_fields = {
            f.field for f in parsed.findings if f.code == "MISSING_METADATA_FIELD"
        }
        assert missing_fields == {
            "dataset_name",
            "experimental_date",
            "organism",
            "experiment_type",
            "measurement_type",
            "scientific_objective",
            "assay",
            "provenance",
            "processing_normalization",
            "time_unit",
        }

    def test_every_measurement_column_is_flagged_as_unit_undeclared(self):
        parsed = parse_workbook(fixture_workbook_bytes())
        undeclared = {
            f.column_name for f in parsed.findings if f.code == "MEASUREMENT_COLUMN_UNIT_UNDECLARED"
        }
        assert undeclared == set(FIXTURE_HEADERS[1:])

    def test_filling_in_dataset_name_removes_that_one_finding_only(self):
        parsed = parse_workbook(
            fixture_workbook_bytes(metadata={"Dataset name": "RAS/MAPK time course"})
        )
        assert parsed.metadata["dataset_name"] == "RAS/MAPK time course"
        missing_fields = {f.field for f in parsed.findings if f.code == "MISSING_METADATA_FIELD"}
        assert "dataset_name" not in missing_fields
        assert "organism" in missing_fields
