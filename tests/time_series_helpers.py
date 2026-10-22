"""Shared workbook-building helpers for the Downloadable Two-Tab Excel Template increment's
own test files (``tests/unit``, ``tests/persistence``, ``tests/api``) -- never duplicated
per file.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.time_series.template import generate_template_workbook

#: The exact regression fixture supplied with this increment's own task instructions -- a
#: real RAS/MAPK phospho-time-course, units/sample identities/replicate assignments/most
#: metadata deliberately left unspecified (task instructions: "retain those gaps honestly").
#: Row 2 (t=120) has a real, reported zero for ``[pMEKtot]`` -- the "preserve reported zeros"
#: regression case.
FIXTURE_HEADERS = ("Time", "[aRAStot]", "[pMEKtot]", "[pYERKtot]", "[pTpYErk]", "[pBRAFtot]")
FIXTURE_ROWS: tuple[tuple[float, ...], ...] = (
    (0, 2.7, 5.85, 1.31, 2.76, 13.95),
    (120, 11.22, 0, 4.72, 9.43, 12.39),
    (240, 21.89, 43.31, 28.45, 54.07, 9.83),
    (480, 21.22, 58.01, 36.68, 66.44, 9.2),
    (720, 18.09, 41.86, 29.54, 54.34, 9.04),
)


def blank_template_bytes() -> bytes:
    return generate_template_workbook()


def populate_data_tab(
    template_bytes: bytes,
    *,
    headers: tuple[str, ...],
    rows: tuple[tuple, ...],
    metadata: dict[str, str] | None = None,
) -> bytes:
    """Starting from a real template's own bytes, overwrite the Data tab's header row with
    ``headers`` and append ``rows`` beneath it, and set any ``metadata`` scalar field (by its
    exact template label) on the Metadata tab. Every field not named in ``metadata`` is left
    exactly as the blank template already has it -- an honest, disclosed gap, never
    fabricated."""
    workbook = load_workbook(BytesIO(template_bytes))
    data_sheet = workbook["Data"]
    for col, value in enumerate(headers, start=1):
        data_sheet.cell(row=1, column=col, value=value)
    for col in range(len(headers) + 1, data_sheet.max_column + 1):
        # openpyxl's cell(..., value=None) is a no-op, not a clear -- direct attribute
        # assignment is the only way to actually blank out a cell that already has a value.
        data_sheet.cell(row=1, column=col).value = None
    for row_offset, row_values in enumerate(rows, start=2):
        for col, value in enumerate(row_values, start=1):
            data_sheet.cell(row=row_offset, column=col, value=value)

    if metadata:
        metadata_sheet = workbook["Metadata"]
        for row_index in range(1, metadata_sheet.max_row + 1):
            label = metadata_sheet.cell(row=row_index, column=1).value
            if label in metadata:
                metadata_sheet.cell(row=row_index, column=2, value=metadata[label])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def fixture_workbook_bytes(metadata: dict[str, str] | None = None) -> bytes:
    """The supplied regression fixture, populated into a real template -- Metadata left blank
    except whatever ``metadata`` overrides explicitly (default: nothing, matching the task's
    own "retain those gaps honestly" instruction for this specific fixture)."""
    return populate_data_tab(
        blank_template_bytes(), headers=FIXTURE_HEADERS, rows=FIXTURE_ROWS, metadata=metadata
    )


def declare_measurement_columns(
    workbook_bytes: bytes, declarations: dict[str, tuple[str | None, str | None]]
) -> bytes:
    """Starting from a workbook's own bytes, write one row into the Metadata tab's own
    "## Measurement column definitions" sub-table per ``declarations`` entry (``column_name ->
    (unit, description)``), found by its own marker text rather than a fixed row number
    (``app.time_series.contract.MEASUREMENT_COLUMNS_SECTION_MARKER``), mirroring the real
    template's own layout exactly (see ``docs/25``'s own "Measurement column definitions"
    sub-table). A column named here with no row at all stays genuinely undeclared -- never
    fabricated."""
    workbook = load_workbook(BytesIO(workbook_bytes))
    metadata_sheet = workbook["Metadata"]
    section_row = next(
        r
        for r in range(1, metadata_sheet.max_row + 1)
        if isinstance(metadata_sheet.cell(row=r, column=1).value, str)
        and metadata_sheet.cell(row=r, column=1).value.startswith(
            "## Measurement column definitions"
        )
    )
    header_row = section_row + 1
    for offset, (column_name, (unit, description)) in enumerate(declarations.items(), start=1):
        metadata_sheet.cell(row=header_row + offset, column=1, value=column_name)
        metadata_sheet.cell(row=header_row + offset, column=2, value=unit)
        metadata_sheet.cell(row=header_row + offset, column=3, value=description)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


#: A second, richer regression fixture for the Read-Only Dataset Inspection and Time-Series
#: Plotting increment (``docs/28``) -- the original RAS/MAPK fixture above deliberately has no
#: sample/condition/replicate columns at all and every measurement column's unit left
#: undeclared; this one exercises every axis that fixture cannot: three distinct samples
#: across two conditions with explicit biological/technical replicate identifiers, one
#: genuinely repeated time (``S3`` at ``t=60``, twice, with different values -- the "repeated
#: times" case, deliberately ambiguous rather than silently averaged/first-wins), rows
#: physically out of time order in the workbook itself (the "unsorted times" case -- row order
#: is never reordered by anything that reads this fixture), and every non-``MEASURED`` cell
#: status at least once (``NA`` -> ``NOT_APPLICABLE``, a genuinely blank cell -> ``MISSING``,
#: ``unknown`` -> ``UNKNOWN``, ``overflow`` -> ``UNRECOGNIZED``). Units are declared for three
#: of its four measurement columns via ``declare_measurement_columns`` (``[aRAStot]``/
#: ``[pBRAFtot]`` both ``"mM"`` -- one shared default display panel; ``[pMEKtot]`` ``"AU"`` --
#: dimensionally incompatible with ``mM``, its own separate panel; ``[undeclared_col]`` left
#: genuinely undeclared -- its own separate "unit not declared" panel).
RICH_FIXTURE_HEADERS = (
    "time",
    "sample_id",
    "condition_id",
    "biological_replicate",
    "technical_replicate",
    "[aRAStot]",
    "[pBRAFtot]",
    "[pMEKtot]",
    "[undeclared_col]",
)
RICH_FIXTURE_ROWS: tuple[tuple, ...] = (
    (0, "S2", "ctrl", "2", "1", 2.1, 10.5, 5.2, 1.0),
    (0, "S1", "ctrl", "1", "1", 2.0, 10.0, 5.0, 1.1),
    (0, "S3", "treated", "1", "1", 2.2, 11.0, 5.5, 0.9),
    (60, "S2", "ctrl", "2", "1", 3.9, 9.3, "overflow", 1.3),
    (60, "S1", "ctrl", "1", "1", 4.0, 9.0, 6.0, 1.2),
    (60, "S3", "treated", "1", "1", 7.0, 7.0, 9.0, 2.0),
    (120, "S1", "ctrl", "1", "1", "NA", 8.0, "unknown", None),
    (60, "S3", "treated", "1", "1", 7.5, 6.8, 9.2, 2.1),
    (120, "S2", "ctrl", "2", "1", 5.0, 8.2, 6.5, 1.4),
    (120, "S3", "treated", "1", "1", 9.0, 6.0, 10.0, 2.2),
)
RICH_FIXTURE_UNIT_DECLARATIONS: dict[str, tuple[str | None, str | None]] = {
    "[aRAStot]": ("mM", "Total RAS"),
    "[pBRAFtot]": ("mM", "Total phospho-BRAF"),
    "[pMEKtot]": ("AU", "Total phospho-MEK, arbitrary units"),
}


def rich_fixture_workbook_bytes(metadata: dict[str, str] | None = None) -> bytes:
    """``RICH_FIXTURE_HEADERS``/``RICH_FIXTURE_ROWS``, with ``RICH_FIXTURE_UNIT_DECLARATIONS``
    written into the Metadata tab's own measurement-column-definitions sub-table."""
    data = populate_data_tab(
        blank_template_bytes(),
        headers=RICH_FIXTURE_HEADERS,
        rows=RICH_FIXTURE_ROWS,
        metadata=metadata,
    )
    return declare_measurement_columns(data, RICH_FIXTURE_UNIT_DECLARATIONS)


__all__ = [
    "FIXTURE_HEADERS",
    "FIXTURE_ROWS",
    "RICH_FIXTURE_HEADERS",
    "RICH_FIXTURE_ROWS",
    "RICH_FIXTURE_UNIT_DECLARATIONS",
    "blank_template_bytes",
    "declare_measurement_columns",
    "fixture_workbook_bytes",
    "populate_data_tab",
    "rich_fixture_workbook_bytes",
]
