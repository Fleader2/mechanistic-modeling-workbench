"""Shared workbook-building helpers for the Initial-Conditions Dataset Import increment's own
test files -- mirrors ``tests.time_series_helpers`` exactly."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.initial_conditions.template import generate_initial_conditions_template_workbook

#: A real, minimal EGFR-Ras-style protein-abundance profile.
FIXTURE_HEADERS = ("target_name", "Abundance")
FIXTURE_ROWS: tuple[tuple, ...] = (
    ("EGFR", 208),
    ("Grb2", 39),
    ("SOS1", 5),
    ("RAS", 98),
)


def blank_template_bytes() -> bytes:
    return generate_initial_conditions_template_workbook()


def populate_data_tab(
    template_bytes: bytes,
    *,
    headers: tuple[str, ...],
    rows: tuple[tuple, ...],
    metadata: dict[str, str] | None = None,
) -> bytes:
    workbook = load_workbook(BytesIO(template_bytes))
    data_sheet = workbook["Data"]
    for col, value in enumerate(headers, start=1):
        data_sheet.cell(row=1, column=col, value=value)
    for col in range(len(headers) + 1, data_sheet.max_column + 1):
        data_sheet.cell(row=1, column=col).value = None
    for row_offset, row_values in enumerate(rows, start=2):
        for col, value in enumerate(row_values, start=1):
            data_sheet.cell(row=row_offset, column=col, value=value)

    if metadata:
        metadata_sheet = workbook["Metadata"]
        for row_index in range(1, metadata_sheet.max_row + 1):
            label = metadata_sheet.cell(row=row_index, column=1).value
            if label in metadata:
                metadata_sheet.cell(row=row_index, column=2, value=metadata[label])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def fixture_workbook_bytes(metadata: dict[str, str] | None = None) -> bytes:
    return populate_data_tab(
        blank_template_bytes(), headers=FIXTURE_HEADERS, rows=FIXTURE_ROWS, metadata=metadata
    )


def declare_measurement_columns(
    workbook_bytes: bytes, declarations: dict[str, tuple[str | None, str | None]]
) -> bytes:
    """Write rows into the Metadata tab's "## Measurement column definitions" sub-table --
    ``declarations`` is ``{column_name: (unit, description)}``."""
    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook["Metadata"]
    marker_row = next(
        r
        for r in range(1, sheet.max_row + 1)
        if isinstance(sheet.cell(row=r, column=1).value, str)
        and sheet.cell(row=r, column=1).value.startswith("## Measurement column definitions")
    )
    row = marker_row + 2  # skip the marker row and the header row beneath it
    for column_name, (unit, description) in declarations.items():
        sheet.cell(row=row, column=1, value=column_name)
        sheet.cell(row=row, column=2, value=unit)
        sheet.cell(row=row, column=3, value=description)
        row += 1
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


__all__ = [
    "FIXTURE_HEADERS",
    "FIXTURE_ROWS",
    "blank_template_bytes",
    "declare_measurement_columns",
    "fixture_workbook_bytes",
    "populate_data_tab",
]
