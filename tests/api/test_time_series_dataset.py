"""Tests for the Downloadable Two-Tab Excel Template increment's own API surface: template
download, preview, save, list, and reopen. No harness/orchestration involvement at all in
this increment -- nothing here is mocked."""

from __future__ import annotations

import uuid
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config.settings import API_V1_PREFIX
from tests.persistence.conftest import make_project_id, make_subproject_id
from tests.time_series_helpers import blank_template_bytes, fixture_workbook_bytes

pytestmark = pytest.mark.api

PROJECTS_URL = f"{API_V1_PREFIX}/projects"

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _base_url(project_id, subproject_id) -> str:
    return f"{PROJECTS_URL}/{project_id}/subprojects/{subproject_id}/datasets/time-series"


class TestTemplateDownload:
    def test_downloads_a_valid_two_tab_workbook(self, client: TestClient, db_session):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()

        response = client.get(f"{_base_url(project_id, subproject_id)}/template")
        assert response.status_code == 200
        assert response.headers["content-type"] == _XLSX_MEDIA_TYPE
        assert "attachment" in response.headers["content-disposition"]

        workbook = load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == ["Metadata", "Data"]

    def test_unknown_project_is_404(self, client: TestClient, db_session):
        subproject_id = make_subproject_id(db_session)
        db_session.commit()
        response = client.get(f"{_base_url(uuid.uuid4(), subproject_id)}/template")
        assert response.status_code == 404


class TestPreview:
    def test_previews_the_supplied_regression_fixture_without_persisting_anything(
        self, client: TestClient, db_session
    ):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()

        response = client.post(
            f"{_base_url(project_id, subproject_id)}/preview",
            files={"file": ("ras_mapk.xlsx", fixture_workbook_bytes(), _XLSX_MEDIA_TYPE)},
        )
        assert response.status_code == 200
        body = response.json()
        assert len(body["rows"]) == 5
        assert {c["column_name"] for c in body["measurement_columns"]} == {
            "[aRAStot]",
            "[pMEKtot]",
            "[pYERKtot]",
            "[pTpYErk]",
            "[pBRAFtot]",
        }
        assert body["rows"][1]["measurements"]["[pMEKtot]"] == {
            "status": "MEASURED",
            "value": 0.0,
            "raw": "0",
        }
        assert any(f["code"] == "MISSING_METADATA_FIELD" for f in body["findings"])

        listing = client.get(f"{_base_url(project_id, subproject_id)}")
        assert listing.json() == []

    def test_a_malformed_workbook_returns_422_with_findings_never_a_500(
        self, client: TestClient, db_session
    ):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()

        response = client.post(
            f"{_base_url(project_id, subproject_id)}/preview",
            files={"file": ("blank.xlsx", blank_template_bytes(), _XLSX_MEDIA_TYPE)},
        )
        assert response.status_code == 422
        body = response.json()
        assert body["detail"]["code"] == "VALIDATION_ERROR"
        assert body["detail"]["context"]["findings"][0]["code"] == "NO_MEASUREMENT_COLUMNS"

    def test_completely_unreadable_bytes_also_returns_422_not_a_500(
        self, client: TestClient, db_session
    ):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()

        response = client.post(
            f"{_base_url(project_id, subproject_id)}/preview",
            files={"file": ("garbage.xlsx", b"not an xlsx file", _XLSX_MEDIA_TYPE)},
        )
        assert response.status_code == 422
        assert response.json()["detail"]["context"]["findings"][0]["code"] == "WORKBOOK_UNREADABLE"


class TestSaveAndReopen:
    def test_save_then_list_then_reopen_round_trips_everything(
        self, client: TestClient, db_session
    ):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()
        base = _base_url(project_id, subproject_id)

        save_response = client.post(
            base,
            files={"file": ("ras_mapk.xlsx", fixture_workbook_bytes(), _XLSX_MEDIA_TYPE)},
            data={"saved_by": "dr-test"},
        )
        assert save_response.status_code == 201
        saved = save_response.json()
        assert saved["saved_by"] == "dr-test"
        assert len(saved["content"]["rows"]) == 5

        listing = client.get(base).json()
        assert len(listing) == 1
        assert listing[0]["id"] == saved["id"]
        assert listing[0]["row_count"] == 5
        assert listing[0]["measurement_column_count"] == 5

        reopened = client.get(f"{base}/{saved['id']}")
        assert reopened.status_code == 200
        assert reopened.json() == saved

    def test_saving_a_malformed_workbook_is_rejected_and_nothing_is_listed(
        self, client: TestClient, db_session
    ):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()
        base = _base_url(project_id, subproject_id)

        response = client.post(
            base, files={"file": ("blank.xlsx", blank_template_bytes(), _XLSX_MEDIA_TYPE)}
        )
        assert response.status_code == 422
        assert client.get(base).json() == []

    def test_resaving_the_identical_workbook_reuses_the_same_dataset(
        self, client: TestClient, db_session
    ):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()
        base = _base_url(project_id, subproject_id)
        data = fixture_workbook_bytes()

        first = client.post(base, files={"file": ("a.xlsx", data, _XLSX_MEDIA_TYPE)})
        second = client.post(base, files={"file": ("a.xlsx", data, _XLSX_MEDIA_TYPE)})
        assert first.status_code == 201
        assert second.status_code == 201
        assert first.json()["id"] == second.json()["id"]
        assert len(client.get(base).json()) == 1

    def test_reopening_an_unknown_dataset_id_is_404(self, client: TestClient, db_session):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()
        response = client.get(f"{_base_url(project_id, subproject_id)}/{uuid.uuid4()}")
        assert response.status_code == 404

    def test_a_dataset_is_not_visible_from_a_different_study(self, client: TestClient, db_session):
        project_id = make_project_id(db_session)
        subproject_a = make_subproject_id(db_session, project_id=project_id)
        subproject_b = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()

        saved = client.post(
            _base_url(project_id, subproject_a),
            files={"file": ("a.xlsx", fixture_workbook_bytes(), _XLSX_MEDIA_TYPE)},
        ).json()

        response = client.get(f"{_base_url(project_id, subproject_b)}/{saved['id']}")
        assert response.status_code == 404
