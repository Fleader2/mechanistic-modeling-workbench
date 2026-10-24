"""Tests for the Initial-Conditions Dataset Import increment's own API surface: template
download, preview, save, list, and reopen. Mirrors
``tests.api.test_time_series_dataset``'s own identical convention -- no harness/orchestration
involvement, nothing mocked.
"""

from __future__ import annotations

import uuid
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config.settings import API_V1_PREFIX
from tests.initial_conditions_helpers import (
    blank_template_bytes,
    fixture_workbook_bytes,
    populate_data_tab,
)
from tests.persistence.conftest import make_project_id, make_subproject_id

pytestmark = pytest.mark.api

PROJECTS_URL = f"{API_V1_PREFIX}/projects"

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _base_url(project_id, subproject_id) -> str:
    return f"{PROJECTS_URL}/{project_id}/subprojects/{subproject_id}/datasets/initial-conditions"


class TestTemplateDownload:
    def test_downloads_a_valid_two_tab_workbook(self, client: TestClient, db_session):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()

        response = client.get(f"{_base_url(project_id, subproject_id)}/template")
        assert response.status_code == 200
        assert response.headers["content-type"] == _XLSX_MEDIA_TYPE
        assert "attachment" in response.headers["content-disposition"]

        workbook = load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == ["Metadata", "Data"]
        header_row = next(workbook["Data"].iter_rows(min_row=1, max_row=1, values_only=True))
        assert header_row == ("target_name", "Abundance")

    def test_unknown_project_is_404(self, client: TestClient, db_session):
        subproject_id = make_subproject_id(db_session)
        db_session.commit()
        response = client.get(f"{_base_url(uuid.uuid4(), subproject_id)}/template")
        assert response.status_code == 404


class TestPreview:
    def test_previews_the_fixture_without_persisting_anything(
        self, client: TestClient, db_session
    ):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()

        response = client.post(
            f"{_base_url(project_id, subproject_id)}/preview",
            files={"file": ("abundance.xlsx", fixture_workbook_bytes(), _XLSX_MEDIA_TYPE)},
        )
        assert response.status_code == 200
        body = response.json()
        assert len(body["rows"]) == 4
        assert {c["column_name"] for c in body["measurement_columns"]} == {"Abundance"}
        assert body["rows"][0]["target_name"] == "EGFR"
        assert body["rows"][0]["measurements"]["Abundance"] == {
            "status": "MEASURED",
            "value": 208.0,
            "raw": "208",
        }
        assert any(f["code"] == "MISSING_METADATA_FIELD" for f in body["findings"])

        listing = client.get(f"{_base_url(project_id, subproject_id)}")
        assert listing.json() == []

    def test_a_malformed_workbook_returns_422_with_findings_never_a_500(
        self, client: TestClient, db_session
    ):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()
        data = populate_data_tab(
            blank_template_bytes(), headers=("Species", "Abundance"), rows=(("EGFR", 208),)
        )

        response = client.post(
            f"{_base_url(project_id, subproject_id)}/preview",
            files={"file": ("bad.xlsx", data, _XLSX_MEDIA_TYPE)},
        )
        assert response.status_code == 422
        assert any(
            f["code"] == "MISSING_TARGET_NAME_COLUMN"
            for f in response.json()["detail"]["context"]["findings"]
        )


class TestSaveAndReopen:
    def test_saves_and_the_save_response_matches_a_reopen(
        self, client: TestClient, db_session
    ):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()

        save_response = client.post(
            f"{_base_url(project_id, subproject_id)}",
            files={"file": ("abundance.xlsx", fixture_workbook_bytes(), _XLSX_MEDIA_TYPE)},
            data={"saved_by": "dr-test"},
        )
        assert save_response.status_code == 201
        saved = save_response.json()
        assert saved["saved_by"] == "dr-test"
        dataset_id = saved["id"]

        reopened = client.get(f"{_base_url(project_id, subproject_id)}/{dataset_id}")
        assert reopened.status_code == 200
        assert reopened.json() == saved

        listing = client.get(f"{_base_url(project_id, subproject_id)}")
        assert listing.status_code == 200
        assert [d["id"] for d in listing.json()] == [dataset_id]

    def test_saving_a_malformed_workbook_is_rejected_and_nothing_is_listed(
        self, client: TestClient, db_session
    ):
        project_id = make_project_id(db_session)
        subproject_id = make_subproject_id(db_session, project_id=project_id)
        db_session.commit()
        data = populate_data_tab(
            blank_template_bytes(), headers=("Species", "Abundance"), rows=(("EGFR", 208),)
        )

        response = client.post(
            f"{_base_url(project_id, subproject_id)}",
            files={"file": ("bad.xlsx", data, _XLSX_MEDIA_TYPE)},
        )
        assert response.status_code == 422
        listing = client.get(f"{_base_url(project_id, subproject_id)}")
        assert listing.json() == []
