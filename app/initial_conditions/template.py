"""Generates the blank, downloadable two-tab initial-conditions dataset workbook (``.xlsx``) --
exactly the layout ``app.initial_conditions.parser`` reads back
(``app.initial_conditions.contract`` is the one shared source of truth for both).

Pure function, no I/O beyond returning bytes -- mirrors ``app.time_series.template`` exactly,
minus the condition-legend and independent-variable-column sub-tables (this dataset type has
neither).
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.initial_conditions.contract import (
    DATA_SHEET_NAME,
    MEASUREMENT_COLUMNS_HEADER,
    MEASUREMENT_COLUMNS_SECTION_MARKER,
    MEASUREMENT_UNCERTAINTY_COLUMNS_HEADER,
    MEASUREMENT_UNCERTAINTY_COLUMNS_SECTION_MARKER,
    METADATA_FIELDS,
    METADATA_SHEET_NAME,
    STRUCTURAL_DATA_COLUMNS,
    TEMPLATE_VERSION,
)

_HEADER_FONT = Font(bold=True)
_SECTION_FONT = Font(bold=True, italic=True)
_EXAMPLE_FONT = Font(italic=True, color="808080")


def _write_metadata_sheet(sheet: Worksheet) -> None:
    sheet.append(("Field", "Value", "Guidance"))
    for cell in sheet[1]:
        cell.font = _HEADER_FONT

    for field in METADATA_FIELDS:
        value = TEMPLATE_VERSION if field.key == "template_version" else None
        sheet.append((field.label, value, field.guidance))

    sheet.append((None, None, None))
    marker_row = sheet.max_row + 1
    sheet.append((MEASUREMENT_COLUMNS_SECTION_MARKER, None, None))
    sheet[f"A{marker_row}"].font = _SECTION_FONT
    sheet.append(MEASUREMENT_COLUMNS_HEADER)
    for cell in sheet[sheet.max_row]:
        cell.font = _HEADER_FONT
    example_row = (
        "e.g. Abundance",
        "nM",
        "e.g. Initial protein/metabolite concentration -- replace or delete this example row.",
    )
    sheet.append(example_row)
    for cell in sheet[sheet.max_row]:
        cell.font = _EXAMPLE_FONT
    for _ in range(6):
        sheet.append((None, None, None))

    sheet.append((None, None, None, None, None))
    marker_row = sheet.max_row + 1
    sheet.append(
        (
            f"{MEASUREMENT_UNCERTAINTY_COLUMNS_SECTION_MARKER} (optional -- only if the Data "
            "tab has a column carrying already-computed, observation-level uncertainty, e.g. "
            "a standard deviation, for one specific measurement column. The statistic column "
            "is required for the declaration to be used: SD, variance, and standard error are "
            "never treated interchangeably)",
            None,
            None,
            None,
            None,
        )
    )
    sheet[f"A{marker_row}"].font = _SECTION_FONT
    sheet.append(MEASUREMENT_UNCERTAINTY_COLUMNS_HEADER)
    for cell in sheet[sheet.max_row]:
        cell.font = _HEADER_FONT
    example_row = (
        "e.g. Abundance__sigma",
        "e.g. Abundance",
        "SD",
        "nM",
        "e.g. Standard deviation across technical replicates, already computed -- replace or "
        "delete this example row.",
    )
    sheet.append(example_row)
    for cell in sheet[sheet.max_row]:
        cell.font = _EXAMPLE_FONT
    for _ in range(4):
        sheet.append((None, None, None, None, None))

    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 70
    sheet.column_dimensions["D"].width = 70
    sheet.column_dimensions["E"].width = 70


def _write_data_sheet(sheet: Worksheet) -> None:
    sheet.append((*STRUCTURAL_DATA_COLUMNS, "Abundance"))
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
    for column_letter in ("A", "B"):
        sheet.column_dimensions[column_letter].width = 24


def generate_initial_conditions_template_workbook() -> bytes:
    """Return the blank, two-tab ``.xlsx`` template as bytes -- ``Metadata`` then ``Data``."""
    workbook = Workbook()
    metadata_sheet = workbook.active
    assert metadata_sheet is not None
    metadata_sheet.title = METADATA_SHEET_NAME
    _write_metadata_sheet(metadata_sheet)

    data_sheet = workbook.create_sheet(DATA_SHEET_NAME)
    _write_data_sheet(data_sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


__all__ = ["generate_initial_conditions_template_workbook"]
