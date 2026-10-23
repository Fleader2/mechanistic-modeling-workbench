"""Parses a time-series dataset workbook (``.xlsx`` bytes) into a ``ParsedWorkbook`` -- never
touching a database, a model, or an Antimony species name (task instructions: "uploading and
saving a dataset must not require a model or matching Antimony species names").

Two distinct outcomes, by design (task instructions: "distinguish malformed workbooks that
cannot be imported from incomplete scientific metadata that can be saved with unresolved
findings"):

* **Malformed** -- the file cannot be opened at all, a required sheet is missing, the Data
  tab has no usable header row, no ``time`` column, zero measurement columns, or a duplicate
  column header. ``parse_workbook`` raises ``WorkbookMalformedError`` (carrying every such
  problem found, not just the first) -- there is no partial result to preview or save.
* **Parseable, possibly scientifically incomplete** -- everything else (a blank required
  metadata field, an undeclared measurement-column unit, an unrecognized cell value, an
  all-missing data row) becomes a non-blocking ``Finding`` on the returned ``ParsedWorkbook``.
  Every finding here is disclosed, never silently dropped, but none of them prevents a save.
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from io import BytesIO
from typing import Literal

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.models.enums import MeasurementValueStatus, UncertaintyStatisticKind
from app.time_series.contract import (
    CONDITION_LEGEND_SECTION_MARKER,
    DATA_SHEET_NAME,
    INDEPENDENT_VARIABLE_COLUMNS_SECTION_MARKER,
    MEASUREMENT_COLUMNS_SECTION_MARKER,
    MEASUREMENT_UNCERTAINTY_COLUMNS_SECTION_MARKER,
    METADATA_FIELDS,
    METADATA_SHEET_NAME,
    NOT_APPLICABLE_TOKENS,
    STRUCTURAL_DATA_COLUMNS,
    TEMPLATE_VERSION,
    TIME_COLUMN,
    UNCERTAINTY_STATISTIC_KIND_TOKENS,
    UNKNOWN_TOKENS,
)

FindingSeverity = Literal["BLOCKING", "INFO"]


@dataclass(frozen=True, slots=True)
class Finding:
    code: str
    severity: FindingSeverity
    message: str
    field: str | None = None
    column_name: str | None = None
    row_index: int | None = None


class WorkbookMalformedError(ValueError):
    """Raised only for a structural problem that leaves nothing parseable -- never for a gap
    in the scientific content itself (see module docstring)."""

    def __init__(self, findings: list[Finding]) -> None:
        self.findings = findings
        super().__init__("; ".join(f.message for f in findings))


@dataclass(frozen=True, slots=True)
class MeasurementValue:
    status: MeasurementValueStatus
    value: float | None
    raw: str | None


@dataclass(frozen=True, slots=True)
class MeasurementColumnDefinition:
    column_name: str
    display_order: int
    unit: str | None
    description: str | None
    declared_in_metadata: bool


@dataclass(frozen=True, slots=True)
class IndependentVariableColumnDefinition:
    """One Data-tab column explicitly declared, in the Metadata tab's own "Independent
    variable definitions" sub-table, as a model *input* (e.g. a dose column) rather than a
    measured output -- docs/29 §4, docs/31. Always ``declared_in_metadata=True``: unlike a
    measurement column (which exists whether or not its unit is declared), a column is only
    ever an independent-variable column *because* this declaration says so -- there is no
    "undeclared independent variable" state, since without this declaration the identical
    column would simply be treated as a measurement column instead (see module docstring)."""

    column_name: str
    display_order: int
    unit: str | None
    description: str | None
    declared_in_metadata: bool = True


@dataclass(frozen=True, slots=True)
class MeasurementUncertaintyColumnDefinition:
    """One Data-tab column explicitly declared, in the Metadata tab's own "Measurement
    uncertainty definitions" sub-table, as already-computed, observation-level uncertainty
    (e.g. a standard deviation) for one specific measurement column -- calibration
    configuration close-out increment (docs/29 §5, docs/32). Always
    ``declared_in_metadata=True``, the identical "only ever exists because it was declared"
    discipline ``IndependentVariableColumnDefinition`` already establishes."""

    column_name: str
    measurement_column_name: str
    #: Required -- which specific statistic this column's own values actually are
    #: (standard deviation / variance / standard error, never left implicit or inferred; see
    #: ``app.models.enums.UncertaintyStatisticKind``). A declaration whose own statistic cell
    #: is blank or unrecognized never reaches this dataclass at all (``parse_workbook``
    #: discloses it as ``MEASUREMENT_UNCERTAINTY_COLUMN_STATISTIC_KIND_UNRECOGNIZED`` and
    #: treats the column as unusable, exactly like a missing/unresolved target).
    statistic_kind: UncertaintyStatisticKind
    display_order: int
    unit: str | None
    description: str | None
    declared_in_metadata: bool = True


@dataclass(frozen=True, slots=True)
class DataRow:
    row_index: int
    sample_id: str | None
    condition_id: str | None
    biological_replicate: str | None
    technical_replicate: str | None
    time_status: MeasurementValueStatus
    time_value: float | None
    time_raw: str | None
    measurements: dict[str, MeasurementValue]
    #: Mirrors ``measurements`` exactly (identical ``MeasurementValue`` shape, identical
    #: never-coerced disclosure of missing/unknown/not-applicable/unrecognized) -- one entry
    #: per declared ``IndependentVariableColumnDefinition``. Empty for any workbook that never
    #: declares the new sub-table, which is exactly what "preserves every existing import
    #: byte-for-byte" requires.
    independent_variables: dict[str, MeasurementValue] = field(default_factory=dict)
    #: Identical shape to ``independent_variables`` -- one entry per declared
    #: ``MeasurementUncertaintyColumnDefinition``, keyed by the **measurement column name it
    #: describes** (never by the uncertainty column's own name), so a consumer never needs to
    #: know which Data-tab column actually carried the value. Empty for any workbook that never
    #: declares the "Measurement uncertainty definitions" sub-table.
    measurement_uncertainty: dict[str, MeasurementValue] = field(default_factory=dict)


@dataclass(frozen=True, slots=True)
class ParsedWorkbook:
    metadata: dict[str, str | None]
    condition_legend: list[tuple[str, str | None]]
    measurement_columns: list[MeasurementColumnDefinition]
    independent_variable_columns: list[IndependentVariableColumnDefinition]
    measurement_uncertainty_columns: list[MeasurementUncertaintyColumnDefinition]
    rows: list[DataRow]
    findings: list[Finding] = field(default_factory=list)


def _classify_cell_value(
    cell_value: object,
) -> tuple[MeasurementValueStatus, float | None, str | None]:
    """One Data-tab measurement/time cell -> ``(status, numeric value, raw text)``. The one
    place "missing," "not applicable," "unknown," "a real zero," and "unparseable junk" are
    told apart (task instructions: "distinguish missing, unknown, and not-applicable values;
    preserve reported zeros").

    A value that parses but is not finite (``inf``/``-inf``/``nan`` -- Python's own ``float()``
    happily parses the literal text "inf"/"nan" into exactly those) is deliberately treated as
    ``UNRECOGNIZED``, never ``MEASURED``: no real physical measurement or uncertainty is ever
    infinite or undefined, and letting either through as a plain float would let a single
    mistyped cell silently produce an infinite or undefined weight downstream (e.g. inverse-
    variance weighting's own ``1/sigma**2``) instead of a disclosed, non-blocking finding.
    """
    if cell_value is None:
        return MeasurementValueStatus.MISSING, None, None
    if isinstance(cell_value, bool):
        # isinstance(True, int) is True in Python -- guarded separately so a stray boolean
        # cell is never silently coerced into 1.0/0.0.
        return MeasurementValueStatus.UNRECOGNIZED, None, str(cell_value)
    if isinstance(cell_value, int | float):
        value = float(cell_value)
        if not math.isfinite(value):
            return MeasurementValueStatus.UNRECOGNIZED, None, str(cell_value)
        return MeasurementValueStatus.MEASURED, value, str(cell_value)
    if isinstance(cell_value, str):
        text = cell_value.strip()
        if text == "":
            return MeasurementValueStatus.MISSING, None, None
        lowered = text.lower()
        if lowered in NOT_APPLICABLE_TOKENS:
            return MeasurementValueStatus.NOT_APPLICABLE, None, text
        if lowered in UNKNOWN_TOKENS:
            return MeasurementValueStatus.UNKNOWN, None, text
        try:
            value = float(text)
        except ValueError:
            return MeasurementValueStatus.UNRECOGNIZED, None, text
        if not math.isfinite(value):
            return MeasurementValueStatus.UNRECOGNIZED, None, text
        return MeasurementValueStatus.MEASURED, value, text
    return MeasurementValueStatus.UNRECOGNIZED, None, str(cell_value)


def _cell_text(value: object) -> str | None:
    """A Metadata/legend-table cell's own value as a trimmed string, or ``None`` for blank --
    never fabricated, never coerced through float formatting that could alter a human-entered
    string (e.g. a date)."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _find_section_start(sheet: Worksheet, marker_prefix: str) -> int | None:
    """The 1-based row number immediately after a section-marker row whose own column-A text
    starts with ``marker_prefix`` -- or ``None`` if no such marker exists (the sub-table is
    simply absent, never an error)."""
    for row_index in range(1, sheet.max_row + 1):
        cell_value = sheet.cell(row=row_index, column=1).value
        if isinstance(cell_value, str) and cell_value.strip().startswith(marker_prefix):
            return row_index + 1
    return None


def _read_subtable(sheet: Worksheet, header_row: int, column_count: int) -> list[tuple[str, ...]]:
    """Every row after ``header_row`` (which is itself skipped) up to the first fully-blank
    row or the end of the sheet -- each as a tuple of ``column_count`` trimmed strings/``None``.
    """
    rows: list[tuple[str, ...]] = []
    for row_index in range(header_row + 1, sheet.max_row + 1):
        values = [sheet.cell(row=row_index, column=c).value for c in range(1, column_count + 1)]
        if all(v is None for v in values):
            break
        rows.append(tuple(_cell_text(v) for v in values))
    return rows


def _parse_metadata_sheet(
    sheet: Worksheet,
) -> tuple[dict[str, str | None], list[tuple[str, str | None]], list[Finding]]:
    findings: list[Finding] = []
    label_to_key = {f.label.strip().lower(): f for f in METADATA_FIELDS}
    found_labels: dict[str, str | None] = {}
    for row_index in range(1, sheet.max_row + 1):
        label_cell = sheet.cell(row=row_index, column=1).value
        if not isinstance(label_cell, str):
            continue
        field_def = label_to_key.get(label_cell.strip().lower())
        if field_def is None:
            continue
        found_labels[field_def.key] = _cell_text(sheet.cell(row=row_index, column=2).value)

    metadata: dict[str, str | None] = {}
    for field_def in METADATA_FIELDS:
        value = found_labels.get(field_def.key)
        metadata[field_def.key] = value
        if field_def.key == "template_version":
            if value != TEMPLATE_VERSION:
                findings.append(
                    Finding(
                        "TEMPLATE_VERSION_MISSING_OR_UNRECOGNIZED",
                        "INFO",
                        f"Template version is {value!r}, expected {TEMPLATE_VERSION!r}. "
                        "Parsed as the current version regardless.",
                        field=field_def.key,
                    )
                )
            continue
        if field_def.required and value is None:
            findings.append(
                Finding(
                    "MISSING_METADATA_FIELD",
                    "INFO",
                    f"{field_def.label!r} was left blank.",
                    field=field_def.key,
                )
            )

    condition_legend: list[tuple[str, str | None]] = []
    legend_start = _find_section_start(sheet, CONDITION_LEGEND_SECTION_MARKER)
    if legend_start is not None:
        for condition_id, description in _read_subtable(sheet, legend_start, 2):
            if condition_id is not None:
                condition_legend.append((condition_id, description))

    return metadata, condition_legend, findings


def _parse_measurement_column_definitions(
    sheet: Worksheet,
) -> dict[str, tuple[str | None, str | None]]:
    """Metadata-declared ``{column_name: (unit, description)}`` -- matched against the Data
    tab's own headers by the caller, never here (this function knows nothing about the Data
    tab)."""
    declarations: dict[str, tuple[str | None, str | None]] = {}
    section_start = _find_section_start(sheet, MEASUREMENT_COLUMNS_SECTION_MARKER)
    if section_start is None:
        return declarations
    for column_name, unit, description in _read_subtable(sheet, section_start, 3):
        if column_name is not None:
            declarations[column_name] = (unit, description)
    return declarations


def _parse_independent_variable_definitions(
    sheet: Worksheet,
) -> dict[str, tuple[str | None, str | None]]:
    """Metadata-declared ``{column_name: (unit, description)}`` for the "Independent variable
    definitions" sub-table -- identical shape and parsing to
    ``_parse_measurement_column_definitions``, a wholly separate sub-table (docs/29 §4). Absent
    entirely for any workbook that never declares one, which is exactly what "preserves every
    existing import byte-for-byte" requires."""
    declarations: dict[str, tuple[str | None, str | None]] = {}
    section_start = _find_section_start(sheet, INDEPENDENT_VARIABLE_COLUMNS_SECTION_MARKER)
    if section_start is None:
        return declarations
    for column_name, unit, description in _read_subtable(sheet, section_start, 3):
        if column_name is not None:
            declarations[column_name] = (unit, description)
    return declarations


def _parse_measurement_uncertainty_definitions(
    sheet: Worksheet,
) -> dict[str, tuple[str | None, str | None, str | None, str | None]]:
    """Metadata-declared ``{uncertainty_column_name: (measurement_column_name,
    statistic_kind_raw, unit, description)}`` for the "Measurement uncertainty definitions"
    sub-table (docs/32) -- ``statistic_kind_raw`` is the literal, as-typed cell text, resolved
    against ``app.time_series.contract.UNCERTAINTY_STATISTIC_KIND_TOKENS`` by the caller, never
    here. Absent entirely for any workbook that never declares one -- "preserves every existing
    import byte-for-byte."""
    declarations: dict[str, tuple[str | None, str | None, str | None, str | None]] = {}
    section_start = _find_section_start(sheet, MEASUREMENT_UNCERTAINTY_COLUMNS_SECTION_MARKER)
    if section_start is None:
        return declarations
    for column_name, measurement_column_name, statistic_kind_raw, unit, description in (
        _read_subtable(sheet, section_start, 5)
    ):
        if column_name is not None:
            declarations[column_name] = (
                measurement_column_name,
                statistic_kind_raw,
                unit,
                description,
            )
    return declarations


def _parse_data_header(
    sheet: Worksheet,
) -> tuple[dict[str, int], list[tuple[int, str]], list[Finding]]:
    """``(structural_column_indices, measurement_columns_by_index, blocking_findings)`` --
    ``structural_column_indices`` maps a structural key (``app.time_series.contract
    .STRUCTURAL_DATA_COLUMNS``) to its 1-based column index; ``measurement_columns_by_index``
    is ``[(index, exact_original_header_text), ...]`` in left-to-right order."""
    findings: list[Finding] = []
    header_values = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))

    structural_by_key: dict[str, int] = {}
    measurement_columns: list[tuple[int, str]] = []
    seen_headers: dict[str, int] = {}

    for zero_based_index, raw_header in enumerate(header_values):
        column_index = zero_based_index + 1
        if raw_header is None or (isinstance(raw_header, str) and raw_header.strip() == ""):
            continue
        header_text = raw_header if isinstance(raw_header, str) else str(raw_header)
        header_text = header_text.strip()
        lowered = header_text.lower()

        structural_match = next((c for c in STRUCTURAL_DATA_COLUMNS if c == lowered), None)
        dedup_key = lowered if structural_match is not None else header_text
        if dedup_key in seen_headers:
            findings.append(
                Finding(
                    "DUPLICATE_DATA_COLUMN_HEADER",
                    "BLOCKING",
                    f"Column header {header_text!r} appears more than once in the Data tab.",
                    column_name=header_text,
                )
            )
            continue
        seen_headers[dedup_key] = column_index

        if structural_match is not None:
            structural_by_key[structural_match] = column_index
        else:
            measurement_columns.append((column_index, header_text))

    if TIME_COLUMN not in structural_by_key:
        findings.append(
            Finding(
                "MISSING_TIME_COLUMN",
                "BLOCKING",
                "The Data tab has no column named 'time' (case-insensitive).",
            )
        )
    if not measurement_columns:
        findings.append(
            Finding(
                "NO_MEASUREMENT_COLUMNS",
                "BLOCKING",
                "The Data tab has no measurement columns -- nothing to import.",
            )
        )
    return structural_by_key, measurement_columns, findings


def parse_workbook(data: bytes) -> ParsedWorkbook:
    """Parse ``.xlsx`` bytes into a ``ParsedWorkbook``. Raises ``WorkbookMalformedError`` (see
    module docstring) for a structural problem; every scientific-content gap instead becomes a
    non-blocking ``Finding`` on the result."""
    try:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise WorkbookMalformedError(
            [
                Finding(
                    "WORKBOOK_UNREADABLE", "BLOCKING", f"Could not open as an .xlsx file: {exc!r}"
                )
            ]
        ) from exc

    missing_sheets = [
        name for name in (METADATA_SHEET_NAME, DATA_SHEET_NAME) if name not in workbook.sheetnames
    ]
    if missing_sheets:
        raise WorkbookMalformedError(
            [
                Finding(
                    "MISSING_REQUIRED_SHEET",
                    "BLOCKING",
                    f"Required sheet {name!r} is missing.",
                )
                for name in missing_sheets
            ]
        )

    data_sheet = workbook[DATA_SHEET_NAME]
    if data_sheet.max_row < 1 or data_sheet.max_column < 1:
        raise WorkbookMalformedError(
            [Finding("DATA_SHEET_EMPTY", "BLOCKING", "The Data tab has no header row.")]
        )

    structural_by_key, measurement_column_headers, header_findings = _parse_data_header(data_sheet)
    blocking = [f for f in header_findings if f.severity == "BLOCKING"]
    if blocking:
        raise WorkbookMalformedError(blocking)

    metadata_sheet = workbook[METADATA_SHEET_NAME]
    metadata, condition_legend, metadata_findings = _parse_metadata_sheet(metadata_sheet)
    declared_columns = _parse_measurement_column_definitions(metadata_sheet)
    #: Joint Calibration Workflow prerequisite increment (docs/29 §4, docs/31). A Data-tab
    #: column named here is excluded from measurement-column processing below entirely -- it
    #: is a model *input*, never an observed output. Empty for any workbook that never
    #: declares this new sub-table, which is exactly what "preserves every existing import
    #: byte-for-byte" requires: every non-structural column then falls through to measurement
    #: processing exactly as it always has.
    declared_independent_variable_columns = _parse_independent_variable_definitions(
        metadata_sheet
    )
    #: Calibration configuration close-out increment (docs/29 §5, docs/32). A Data-tab column
    #: named here carries already-computed, observation-level uncertainty for one specific
    #: measurement column -- excluded from measurement-column processing below, identically to
    #: an independent-variable column. Empty for any workbook that never declares this new
    #: sub-table.
    declared_uncertainty_columns = _parse_measurement_uncertainty_definitions(metadata_sheet)

    findings: list[Finding] = list(metadata_findings)

    # A column name declared in *both* sub-tables is ambiguous -- the independent-variable
    # declaration wins (it was introduced first, docs/31), and the conflict is disclosed rather
    # than silently resolved.
    conflicting_declarations = set(declared_independent_variable_columns) & set(
        declared_uncertainty_columns
    )
    for name in sorted(conflicting_declarations):
        findings.append(
            Finding(
                "COLUMN_DECLARED_BOTH_INDEPENDENT_VARIABLE_AND_UNCERTAINTY",
                "INFO",
                f"Column {name!r} is declared both as an independent variable and as a "
                "measurement-uncertainty column -- treated as an independent variable.",
                column_name=name,
            )
        )
    declared_uncertainty_columns = {
        name: value
        for name, value in declared_uncertainty_columns.items()
        if name not in conflicting_declarations
    }

    independent_variable_column_headers = [
        (index, name)
        for index, name in measurement_column_headers
        if name in declared_independent_variable_columns
    ]
    uncertainty_column_headers = [
        (index, name)
        for index, name in measurement_column_headers
        if name in declared_uncertainty_columns
    ]
    true_measurement_column_headers = [
        (index, name)
        for index, name in measurement_column_headers
        if name not in declared_independent_variable_columns
        and name not in declared_uncertainty_columns
    ]

    measurement_columns: list[MeasurementColumnDefinition] = []
    declared_names_used: set[str] = set()
    for display_order, (_, column_name) in enumerate(true_measurement_column_headers):
        declaration = declared_columns.get(column_name)
        if declaration is None:
            findings.append(
                Finding(
                    "MEASUREMENT_COLUMN_UNIT_UNDECLARED",
                    "INFO",
                    f"Data column {column_name!r} has no matching entry in the Metadata "
                    "tab's measurement column definitions -- its unit is unknown.",
                    column_name=column_name,
                )
            )
            measurement_columns.append(
                MeasurementColumnDefinition(column_name, display_order, None, None, False)
            )
        else:
            declared_names_used.add(column_name)
            unit, description = declaration
            measurement_columns.append(
                MeasurementColumnDefinition(column_name, display_order, unit, description, True)
            )
    for declared_name in declared_columns:
        if declared_name not in declared_names_used:
            findings.append(
                Finding(
                    "UNDECLARED_MEASUREMENT_COLUMN_ORPHAN_DEFINITION",
                    "INFO",
                    f"The Metadata tab declares a measurement column {declared_name!r} that "
                    "does not appear in the Data tab.",
                    column_name=declared_name,
                )
            )

    independent_variable_columns: list[IndependentVariableColumnDefinition] = []
    independent_names_present: set[str] = set()
    for display_order, (_, column_name) in enumerate(independent_variable_column_headers):
        unit, description = declared_independent_variable_columns[column_name]
        independent_variable_columns.append(
            IndependentVariableColumnDefinition(column_name, display_order, unit, description)
        )
        independent_names_present.add(column_name)
    for declared_name in declared_independent_variable_columns:
        if declared_name not in independent_names_present:
            findings.append(
                Finding(
                    "ORPHAN_INDEPENDENT_VARIABLE_COLUMN_DEFINITION",
                    "INFO",
                    f"The Metadata tab declares an independent variable column "
                    f"{declared_name!r} that does not appear in the Data tab.",
                    column_name=declared_name,
                )
            )

    # ``{uncertainty_column_index: measurement_column_name}`` -- only for a declaration that
    # actually resolves to a real, present measurement column, with no earlier-declared
    # uncertainty column already claiming the same target (first declaration wins; every later
    # one is disclosed as a duplicate, never silently applied on top of or instead of the
    # first). Built before the main per-row loop so that loop stays a single pass.
    true_measurement_names_present = {name for _, name in true_measurement_column_headers}
    measurement_uncertainty_columns: list[MeasurementUncertaintyColumnDefinition] = []
    uncertainty_index_to_target: dict[int, str] = {}
    targets_claimed: set[str] = set()
    uncertainty_names_present: set[str] = set()
    for display_order, (column_index, column_name) in enumerate(uncertainty_column_headers):
        uncertainty_names_present.add(column_name)
        measurement_column_name, statistic_kind_raw, unit, description = (
            declared_uncertainty_columns[column_name]
        )
        if not measurement_column_name:
            findings.append(
                Finding(
                    "MEASUREMENT_UNCERTAINTY_COLUMN_MISSING_TARGET",
                    "INFO",
                    f"Measurement-uncertainty column {column_name!r} does not name which "
                    "measurement column it describes -- its values are not used.",
                    column_name=column_name,
                )
            )
            continue
        if measurement_column_name not in true_measurement_names_present:
            findings.append(
                Finding(
                    "MEASUREMENT_UNCERTAINTY_COLUMN_TARGET_NOT_FOUND",
                    "INFO",
                    f"Measurement-uncertainty column {column_name!r} names "
                    f"{measurement_column_name!r} as the measurement column it describes, but "
                    "no such measurement column exists -- its values are not used.",
                    column_name=column_name,
                )
            )
            continue
        if measurement_column_name in targets_claimed:
            findings.append(
                Finding(
                    "DUPLICATE_MEASUREMENT_UNCERTAINTY_TARGET",
                    "INFO",
                    f"Measurement-uncertainty column {column_name!r} also names "
                    f"{measurement_column_name!r}, already described by an earlier-declared "
                    "uncertainty column -- only the first declaration is used.",
                    column_name=column_name,
                )
            )
            continue
        statistic_kind = UNCERTAINTY_STATISTIC_KIND_TOKENS.get(
            (statistic_kind_raw or "").strip().lower()
        )
        if statistic_kind is None:
            findings.append(
                Finding(
                    "MEASUREMENT_UNCERTAINTY_COLUMN_STATISTIC_KIND_UNRECOGNIZED",
                    "INFO",
                    f"Measurement-uncertainty column {column_name!r} does not declare a "
                    "recognized statistic (standard deviation / variance / standard error) "
                    "-- its values are not used. A standard deviation, a variance, and a "
                    "standard error are not interchangeable, so this is never guessed.",
                    column_name=column_name,
                )
            )
            continue
        targets_claimed.add(measurement_column_name)
        uncertainty_index_to_target[column_index] = measurement_column_name
        measurement_uncertainty_columns.append(
            MeasurementUncertaintyColumnDefinition(
                column_name,
                measurement_column_name,
                statistic_kind,
                display_order,
                unit,
                description,
            )
        )
    for declared_name in declared_uncertainty_columns:
        if declared_name not in uncertainty_names_present:
            findings.append(
                Finding(
                    "ORPHAN_MEASUREMENT_UNCERTAINTY_COLUMN_DEFINITION",
                    "INFO",
                    f"The Metadata tab declares a measurement-uncertainty column "
                    f"{declared_name!r} that does not appear in the Data tab.",
                    column_name=declared_name,
                )
            )

    rows: list[DataRow] = []
    kept_row_index = 0
    for sheet_row_values in data_sheet.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in sheet_row_values):
            continue
        row_index = kept_row_index
        kept_row_index += 1

        def _structural(key: str, values: tuple = sheet_row_values) -> str | None:
            col = structural_by_key.get(key)
            if col is None:
                return None
            return _cell_text(values[col - 1])

        time_col = structural_by_key[TIME_COLUMN]
        time_status, time_value, time_raw = _classify_cell_value(sheet_row_values[time_col - 1])
        if time_status == MeasurementValueStatus.MISSING:
            findings.append(
                Finding("TIME_VALUE_MISSING", "INFO", "Time value is blank.", row_index=row_index)
            )
        elif time_status == MeasurementValueStatus.UNRECOGNIZED:
            findings.append(
                Finding(
                    "DATA_CELL_UNRECOGNIZED_VALUE",
                    "INFO",
                    f"Time value {time_raw!r} is not numeric and not a recognized token.",
                    column_name="time",
                    row_index=row_index,
                )
            )

        measurements: dict[str, MeasurementValue] = {}
        any_measured_or_disclosed = False
        for column_index, column_name in true_measurement_column_headers:
            status, value, raw = _classify_cell_value(sheet_row_values[column_index - 1])
            measurements[column_name] = MeasurementValue(status, value, raw)
            if status != MeasurementValueStatus.MISSING:
                any_measured_or_disclosed = True
            if status == MeasurementValueStatus.UNRECOGNIZED:
                findings.append(
                    Finding(
                        "DATA_CELL_UNRECOGNIZED_VALUE",
                        "INFO",
                        f"Value {raw!r} in column {column_name!r} is not numeric and not a "
                        "recognized token.",
                        column_name=column_name,
                        row_index=row_index,
                    )
                )
        if not any_measured_or_disclosed:
            findings.append(
                Finding(
                    "DATA_ROW_ALL_MEASUREMENTS_MISSING",
                    "INFO",
                    "Every measurement column in this row is blank.",
                    row_index=row_index,
                )
            )

        independent_variables: dict[str, MeasurementValue] = {}
        for column_index, column_name in independent_variable_column_headers:
            status, value, raw = _classify_cell_value(sheet_row_values[column_index - 1])
            independent_variables[column_name] = MeasurementValue(status, value, raw)
            if status == MeasurementValueStatus.UNRECOGNIZED:
                findings.append(
                    Finding(
                        "DATA_CELL_UNRECOGNIZED_VALUE",
                        "INFO",
                        f"Value {raw!r} in independent variable column {column_name!r} is not "
                        "numeric and not a recognized token.",
                        column_name=column_name,
                        row_index=row_index,
                    )
                )

        measurement_uncertainty: dict[str, MeasurementValue] = {}
        for column_index, measurement_column_name in uncertainty_index_to_target.items():
            status, value, raw = _classify_cell_value(sheet_row_values[column_index - 1])
            measurement_uncertainty[measurement_column_name] = MeasurementValue(status, value, raw)
            if status == MeasurementValueStatus.UNRECOGNIZED:
                findings.append(
                    Finding(
                        "DATA_CELL_UNRECOGNIZED_VALUE",
                        "INFO",
                        f"Value {raw!r} in uncertainty column for {measurement_column_name!r} "
                        "is not numeric and not a recognized token.",
                        column_name=measurement_column_name,
                        row_index=row_index,
                    )
                )
            elif status == MeasurementValueStatus.MEASURED and value is not None and value < 0:
                findings.append(
                    Finding(
                        "NEGATIVE_UNCERTAINTY_VALUE",
                        "INFO",
                        f"Uncertainty value {raw!r} for {measurement_column_name!r} is "
                        "negative -- uncertainty is never negative; preserved as reported, "
                        "never flipped or discarded.",
                        column_name=measurement_column_name,
                        row_index=row_index,
                    )
                )

        rows.append(
            DataRow(
                row_index=row_index,
                sample_id=_structural("sample_id"),
                condition_id=_structural("condition_id"),
                biological_replicate=_structural("biological_replicate"),
                technical_replicate=_structural("technical_replicate"),
                time_status=time_status,
                time_value=time_value,
                time_raw=time_raw,
                measurements=measurements,
                independent_variables=independent_variables,
                measurement_uncertainty=measurement_uncertainty,
            )
        )

    return ParsedWorkbook(
        metadata=metadata,
        condition_legend=condition_legend,
        measurement_columns=measurement_columns,
        independent_variable_columns=independent_variable_columns,
        measurement_uncertainty_columns=measurement_uncertainty_columns,
        rows=rows,
        findings=findings,
    )


__all__ = [
    "DataRow",
    "Finding",
    "IndependentVariableColumnDefinition",
    "MeasurementColumnDefinition",
    "MeasurementUncertaintyColumnDefinition",
    "MeasurementValue",
    "ParsedWorkbook",
    "WorkbookMalformedError",
    "parse_workbook",
]
