"""Generates the blank, downloadable two-tab time-series dataset workbook (``.xlsx``) --
exactly the layout ``app.time_series.parser`` reads back (``app.time_series.contract`` is the
one shared source of truth for both).

Pure function, no I/O beyond returning bytes -- the caller (the API route) decides whether that
is a file response, a test assertion, or anything else.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.time_series.contract import (
    CONDITION_LEGEND_HEADER,
    CONDITION_LEGEND_SECTION_MARKER,
    DATA_SHEET_NAME,
    INDEPENDENT_VARIABLE_COLUMNS_HEADER,
    INDEPENDENT_VARIABLE_COLUMNS_SECTION_MARKER,
    MEASUREMENT_COLUMNS_HEADER,
    MEASUREMENT_COLUMNS_SECTION_MARKER,
    MEASUREMENT_UNCERTAINTY_COLUMNS_HEADER,
    MEASUREMENT_UNCERTAINTY_COLUMNS_SECTION_MARKER,
    METADATA_FIELDS,
    METADATA_SHEET_NAME,
    STRUCTURAL_DATA_COLUMNS,
    TEMPLATE_VERSION,
)

_HEADER_FONT = Font(bold=True)
_SECTION_FONT = Font(bold=True, italic=True)
_EXAMPLE_FONT = Font(italic=True, color="808080")


def _write_metadata_sheet(sheet: Worksheet) -> None:
    sheet.append(("Field", "Value", "Guidance"))
    for cell in sheet[1]:
        cell.font = _HEADER_FONT

    for field in METADATA_FIELDS:
        value = TEMPLATE_VERSION if field.key == "template_version" else None
        sheet.append((field.label, value, field.guidance))

    sheet.append((None, None, None))
    marker_row = sheet.max_row + 1
    sheet.append((MEASUREMENT_COLUMNS_SECTION_MARKER, None, None))
    sheet[f"A{marker_row}"].font = _SECTION_FONT
    sheet.append(MEASUREMENT_COLUMNS_HEADER)
    for cell in sheet[sheet.max_row]:
        cell.font = _HEADER_FONT
    example_row = (
        "e.g. glucose_mM",
        "mM",
        "e.g. Extracellular glucose concentration -- replace or delete this example row.",
    )
    sheet.append(example_row)
    for cell in sheet[sheet.max_row]:
        cell.font = _EXAMPLE_FONT
    for _ in range(6):
        sheet.append((None, None, None))

    sheet.append((None, None, None))
    marker_row = sheet.max_row + 1
    sheet.append(
        (
            f"{CONDITION_LEGEND_SECTION_MARKER} (optional -- only if the Data tab uses a "
            "Condition ID column)",
            None,
            None,
        )
    )
    sheet[f"A{marker_row}"].font = _SECTION_FONT
    sheet.append(CONDITION_LEGEND_HEADER)
    for cell in sheet[sheet.max_row]:
        cell.font = _HEADER_FONT
    for _ in range(4):
        sheet.append((None, None))

    sheet.append((None, None, None))
    marker_row = sheet.max_row + 1
    sheet.append(
        (
            f"{INDEPENDENT_VARIABLE_COLUMNS_SECTION_MARKER} (optional -- only if the Data tab "
            "has a column that is a model INPUT, e.g. dose, never an observed output)",
            None,
            None,
        )
    )
    sheet[f"A{marker_row}"].font = _SECTION_FONT
    sheet.append(INDEPENDENT_VARIABLE_COLUMNS_HEADER)
    for cell in sheet[sheet.max_row]:
        cell.font = _HEADER_FONT
    example_row = (
        "e.g. dose_uM",
        "uM",
        "e.g. Drug concentration applied at t=0 -- replace or delete this example row.",
    )
    sheet.append(example_row)
    for cell in sheet[sheet.max_row]:
        cell.font = _EXAMPLE_FONT
    for _ in range(4):
        sheet.append((None, None, None))

    sheet.append((None, None, None, None, None))
    marker_row = sheet.max_row + 1
    sheet.append(
        (
            f"{MEASUREMENT_UNCERTAINTY_COLUMNS_SECTION_MARKER} (optional -- only if the Data "
            "tab has a column carrying already-computed, observation-level uncertainty, e.g. "
            "a standard deviation, for one specific measurement column. The statistic column "
            "is required for the declaration to be used: SD, variance, and standard error are "
            "never treated interchangeably)",
            None,
            None,
            None,
            None,
        )
    )
    sheet[f"A{marker_row}"].font = _SECTION_FONT
    sheet.append(MEASUREMENT_UNCERTAINTY_COLUMNS_HEADER)
    for cell in sheet[sheet.max_row]:
        cell.font = _HEADER_FONT
    example_row = (
        "e.g. glucose_mM__sigma",
        "e.g. glucose_mM",
        "SD",
        "mM",
        "e.g. Standard deviation across technical replicates, already computed -- replace or "
        "delete this example row.",
    )
    sheet.append(example_row)
    for cell in sheet[sheet.max_row]:
        cell.font = _EXAMPLE_FONT
    for _ in range(4):
        sheet.append((None, None, None, None, None))

    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 70
    sheet.column_dimensions["E"].width = 70
    sheet.column_dimensions["D"].width = 70


def _write_data_sheet(sheet: Worksheet) -> None:
    sheet.append(STRUCTURAL_DATA_COLUMNS)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
    for column_letter in ("A", "B", "C", "D", "E"):
        sheet.column_dimensions[column_letter].width = 20


def generate_template_workbook() -> bytes:
    """Return the blank, two-tab ``.xlsx`` template as bytes -- ``Metadata`` then ``Data``, in
    that order (the order a human opening the file sees first)."""
    workbook = Workbook()
    metadata_sheet = workbook.active
    assert metadata_sheet is not None
    metadata_sheet.title = METADATA_SHEET_NAME
    _write_metadata_sheet(metadata_sheet)

    data_sheet = workbook.create_sheet(DATA_SHEET_NAME)
    _write_data_sheet(data_sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


__all__ = ["generate_template_workbook"]
